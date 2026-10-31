"""Load operational limits from the organiser workbook (data/raw/challenge2/challenge-2-operational-constraints.xlsx).

Organiser rule: "Use the limits in this workbook. Any limit you change must be stated in your demo."
So limits are read from the workbook, never hardcoded. Team assumptions layered on top are listed in
ASSUMPTIONS and must appear in the demo / decision document.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
WORKBOOK = REPO / "data" / "raw" / "challenge2" / "challenge-2-operational-constraints.xlsx"


def _first_number(text) -> float:
    m = re.search(r"\d[\d,]*\.?\d*", str(text))
    if not m:
        raise ValueError(f"no number in workbook value {text!r}")
    return float(m.group(0).replace(",", ""))


@dataclass(frozen=True)
class FuelCrewRules:
    final_reserve_min: float          # never plan below this (minutes of holding fuel)
    contingency_pct: float            # % of trip fuel
    burn_narrowbody_kg_h: float
    burn_widebody_kg_h: float
    spare_on_arrival_min: float       # minutes of holding fuel available on arrival
    flight_crew_max_duty_h: float     # from report time
    ground_crew_shift_h: float
    gate_change_notice_min: float
    source: str = str(WORKBOOK.relative_to(REPO))


@dataclass(frozen=True)
class AirportLimits:
    iata: str
    size_class: str
    active_runways: int
    gates: int
    max_runway_ops_per_hour: int
    min_turnaround_min: int
    spare_gates_at_any_time: int


# Team assumptions on top of the workbook. Every one must be stated in the demo.
ASSUMPTIONS = {
    "contingency_inside_spare": "The 5% contingency fuel is counted inside the 45-min spare fuel, so "
                                "holdable fuel = spare fuel; the 30-min final reserve is never used.",
    "cruise_nm_per_min": 7.0,       # ~420 kt groundspeed to the alternate
    "alternate_approach_min": 10.0,  # approach and landing allowance at the alternate
    "crew_report_before_first_leg_min": 60.0,
    "widebody_if_distance_mi_at_least": 2500.0,
    "minimum_fuel_margin_min": 10.0,  # CLAUDE.md pilot red line
    "mct_min": {"Large hub": 45.0, "default": 35.0},
    "connection_window_min": (35.0, 180.0),
    "connection_pax_range": (3, 25),
}


@lru_cache(maxsize=1)
def fuel_crew_rules(path: Path = WORKBOOK) -> FuelCrewRules:
    ws = load_workbook(path, data_only=True)["fuel_and_crew"]
    rows = {str(r[0]).strip(): r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}

    def get(prefix):
        for k, v in rows.items():
            if k.lower().startswith(prefix.lower()):
                return _first_number(v)
        raise KeyError(f"workbook fuel_and_crew sheet has no rule starting with {prefix!r}")

    return FuelCrewRules(
        final_reserve_min=get("Final fuel reserve"),
        contingency_pct=get("Contingency fuel"),
        burn_narrowbody_kg_h=get("Holding fuel burn (narrowbody)"),
        burn_widebody_kg_h=get("Holding fuel burn (widebody)"),
        spare_on_arrival_min=get("Assumed spare fuel"),
        flight_crew_max_duty_h=get("Flight crew max duty"),
        ground_crew_shift_h=get("Ground crew shift"),
        gate_change_notice_min=get("Gate change notice"),
    )


@lru_cache(maxsize=1)
def airport_limits(path: Path = WORKBOOK) -> dict[str, AirportLimits]:
    ws = load_workbook(path, data_only=True)["airport_capacity"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "IATA_CODE")
    hdr = list(rows[hdr_i])
    col = {name: hdr.index(name) for name in hdr if name}
    out = {}
    for r in rows[hdr_i + 1:]:
        if not r or not r[0]:
            continue
        out[r[0]] = AirportLimits(
            iata=r[0], size_class=r[col["size_class"]], active_runways=int(r[col["active_runways"]]),
            gates=int(r[col["gates"]]), max_runway_ops_per_hour=int(r[col["max_runway_ops_per_hour"]]),
            min_turnaround_min=int(r[col["min_turnaround_min"]]),
            spare_gates_at_any_time=int(r[col["spare_gates_at_any_time"]]),
        )
    return out
